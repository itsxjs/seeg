from docx import Document

report = '/Users/defanive/Desktop/开题报告.docx'
form = '/Users/defanive/Desktop/Diploma/心理系2026届本科论文工作指导文件/心理系2026届本科论文工作指导文件/附件8：心理系毕业论文（设计）工作中期检查表.docx'

for p in [report, form]:
    print('\n' + '=' * 40)
    print(p)
    d = Document(p)
    print('paragraph_count', len(d.paragraphs), 'table_count', len(d.tables))
    print('\n[Non-empty paragraphs sample]')
    shown = 0
    for i, para in enumerate(d.paragraphs, 1):
        t = para.text.strip()
        if t:
            print(f'{i:03d}: {t}')
            shown += 1
        if shown >= 80:
            break

    print('\n[Tables]')
    for ti, tab in enumerate(d.tables, 1):
        print(f'Table {ti}: rows={len(tab.rows)}, cols={len(tab.columns)}')
        for r, row in enumerate(tab.rows, 1):
            cells = [c.text.strip().replace('\n', ' / ') for c in row.cells]
            if any(cells):
                print(f'  r{r}: ' + ' || '.join(cells))
        print('-' * 20)
